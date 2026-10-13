from docx import Document
from docx.shared import Pt

def save_docx(topic, report):
    doc = Document()
    title = doc.add_heading(topic, level=1)
    title.runs[0].font.size = Pt(24)

    # Add report content line by line
    for line in report.split("\n"):
        line = line.strip()

        # Markdown headings support
        if line.startswith("# "):
            doc.add_heading(line.replace("# ", ""), level=1)
        elif line.startswith("## "):
            doc.add_heading(line.replace("## ", ""), level=2)
        elif line.startswith("### "):
            doc.add_heading(line.replace("### ", ""), level=3)
        elif line != "":
            doc.add_paragraph(line)

    output_path = f"reports/{topic}.docx"
    doc.save(output_path)
    return output_path